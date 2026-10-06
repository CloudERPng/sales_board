"""Monthly management pack: every Sales Board and Inventory Board report for a closed month in one workbook."""
import io

import frappe
from frappe.utils import add_days, cint, flt, format_date, get_first_day, get_last_day, getdate, now_datetime, today

GREEN, AMBER, TEAL, MINT, CREAM = "1F5F3F", "E9A23B", "2E8B6E", "EAF3EC", "FFF6E5"
BASE = "sales_board.sales_board.report"
FMT = {"Currency": '₦#,##0;(₦#,##0);"-"', "Int": '#,##0;(#,##0);"-"', "Float": '#,##0.0;(#,##0.0);"-"',
	"Percent": '0.0"%";(0.0"%");"-"', "Date": "dd-mmm-yy"}

# (sheet, report module, extra filters, uses period or as-of)
SECTIONS = [
	("Sales", [
		("Order Funnel", "order_funnel", {}, "period"),
		("Revenue Recon", "revenue_reconciliation", {}, "period"),
		("Marketing ROAS", "marketing_roas", {"view_by": "Marketer × Product"}, "period"),
		("Marketers", "marketer_scorecard", {}, "period"),
		("Closers", "closer_scorecard", {}, "period"),
		("Delivery Agents", "delivery_agent_scorecard", {}, "period"),
		("Agent Remittance", "agent_remittance", {}, "period"),
		("Pipeline by Closer", "pipeline_recovery", {"view_by": "Closer"}, "period"),
		("By Product", "sales_performance_by_dimension", {"group_by": "Product"}, "period"),
		("By Zone", "sales_performance_by_dimension", {"group_by": "Zone"}, "period"),
		("By State", "sales_performance_by_dimension", {"group_by": "State"}, "period"),
		("Daily Trend", "sales_performance_by_dimension", {"group_by": "Order Date"}, "period"),
		("By Hour", "sales_performance_by_dimension", {"group_by": "Hour of Day"}, "period"),
	]),
	("Inventory", [
		("Stock by Location", "stock_by_location", {"view_by": "Location"}, "asof"),
		("Agent Stock Statement", "agent_stock_statement", {"view_by": "Location", "location_type": "Agent"}, "period"),
		("Days of Cover", "days_of_cover", {"view_by": "Location", "location_type": "Agent"}, "asof"),
		("Rebalancing", "days_of_cover", {"view_by": "Rebalancing suggestions", "location_type": "Agent"}, "asof"),
		("Distribution", "distribution_efficiency", {"view_by": "Location", "location_type": "Agent"}, "period"),
		("Stock vs Demand", "distribution_efficiency", {"view_by": "Zone", "location_type": "Agent"}, "period"),
		("Lost to No Stock", "lost_sales_from_stock_outs", {"view_by": "Zone"}, "period"),
		("Idle Stock", "idle_stock", {}, "asof"),
		("Transfers", "stock_transfers", {"view_by": "Route"}, "period"),
		("Adjustments", "inventory_adjustments", {"view_by": "Location"}, "period"),
		("Adjustment Detail", "inventory_adjustments", {"view_by": "Detail"}, "period"),
	]),
]


def run_report(module, company, start, end, extra, kind):
	f = {"company": company, **extra}
	if kind == "asof":
		f["as_of"] = str(end)
	else:
		f["from_date"], f["to_date"] = str(start), str(end)
	res = frappe.get_attr(f"{BASE}.{module}.{module}.execute")(f)
	cols, data = res[0], res[1]
	return frappe._dict(columns=cols, data=data, message=res[2] if len(res) > 2 else None,
		chart=res[3] if len(res) > 3 else None, summary=res[4] if len(res) > 4 else None)


def build_workbook(company, start, end, commentary=None):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter as CL

	F = "Arial"
	thin = Side(style="thin", color="C9D6CC")
	wb = Workbook()
	summary_ws = wb.active
	summary_ws.title = "Summary"
	results, errors = {}, []
	for section, items in SECTIONS:
		for sheet, module, extra, kind in items:
			try:
				results[sheet] = (section, run_report(module, company, start, end, extra, kind))
			except Exception:
				errors.append(sheet)
				frappe.log_error(title=f"Sales Board pack: {sheet}")

	def head(ws, text, sub):
		ws["A1"] = text
		ws["A1"].font = Font(name=F, size=15, bold=True, color=GREEN)
		ws["A2"] = sub
		ws["A2"].font = Font(name=F, size=9, italic=True, color="6B6B6B")
		ws.sheet_view.showGridLines = False

	period = f"{format_date(start)} to {format_date(end)}"
	for section, items in SECTIONS:
		for sheet, *_x in items:
			if sheet not in results:
				continue
			_s, r = results[sheet]
			ws = wb.create_sheet(sheet)
			ws.sheet_properties.tabColor = GREEN if section == "Sales" else AMBER
			head(ws, sheet, f"{company} | {period} | {section}")
			row = 3
			if r.message:
				ws.cell(row, 1, frappe.utils.strip_html(r.message)).font = Font(name=F, size=9, italic=True, color="6B6B6B")
				row += 1
			row += 1
			cols = [c for c in r.columns if isinstance(c, dict)]
			for i, c in enumerate(cols):
				x = ws.cell(row, i + 1, c.get("label"))
				x.font = Font(name=F, bold=True, color="FFFFFF", size=10)
				x.fill = PatternFill("solid", fgColor=GREEN)
				x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
				ws.column_dimensions[CL(i + 1)].width = max(10, min(40, cint(c.get("width") or 110) / 7))
			ws.row_dimensions[row].height = 36
			hdr_row = row
			for j, d in enumerate(r.data):
				row += 1
				bold = bool(d.get("bold"))
				for i, c in enumerate(cols):
					v = d.get(c["fieldname"])
					if isinstance(v, (list, tuple, set, dict)):
						v = str(v)
					x = ws.cell(row, i + 1, v)
					x.number_format = FMT.get(c.get("fieldtype"), "General")
					x.font = Font(name=F, size=10, bold=bold)
					x.border = Border(bottom=thin)
					if bold:
						x.fill = PatternFill("solid", fgColor=CREAM)
					elif j % 2:
						x.fill = PatternFill("solid", fgColor=MINT)
			ws.freeze_panes = ws.cell(hdr_row + 1, 2)
			if r.data:
				ws.auto_filter.ref = f"A{hdr_row}:{CL(len(cols))}{row}"
			add_chart(ws, r.chart, row + 3, len(cols) + 2)

	# ---------------- summary
	ws = summary_ws
	head(ws, f"{company} — Monthly Management Pack", f"{period} | Frozen month-end figures where available | "
		f"Generated {frappe.utils.format_datetime(now_datetime())}")
	ws.column_dimensions["A"].width = 38
	ws.column_dimensions["B"].width = 22
	row = 4
	if commentary:
		ws.cell(row, 1, "Management commentary").font = Font(name=F, size=12, bold=True, color=TEAL)
		row += 1
		for para in [p.strip() for p in frappe.utils.html2text(commentary).split("\n") if p.strip()]:
			ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
			x = ws.cell(row, 1, para)
			x.font = Font(name=F, size=10)
			x.alignment = Alignment(wrap_text=True, vertical="top")
			ws.row_dimensions[row].height = max(15, 15 * (len(para) // 110 + 1))
			row += 1
		row += 1
	for title, sheets in [("Sales", ["Order Funnel", "Marketing ROAS", "Agent Remittance"]),
		("Inventory", ["Stock by Location", "Agent Stock Statement", "Days of Cover", "Lost to No Stock", "Adjustments"])]:
		ws.cell(row, 1, title).font = Font(name=F, size=12, bold=True, color=TEAL)
		row += 1
		seen = {"Delivered value"}
		for sheet in sheets:
			r = results.get(sheet, (None, None))[1]
			for k in (r.summary or []) if r else []:
				if k["label"] in seen:
					continue
				seen.add(k["label"])
				ws.cell(row, 1, k["label"]).font = Font(name=F, size=10)
				x = ws.cell(row, 2, k["value"])
				x.number_format = FMT.get(k.get("datatype"), "#,##0.00")
				x.font = Font(name=F, size=10, bold=True, color=GREEN)
				ws.cell(row, 1).border = Border(bottom=thin)
				x.border = Border(bottom=thin)
				row += 1
		row += 1
	ws.cell(row, 1, "Contents").font = Font(name=F, size=12, bold=True, color=TEAL)
	row += 1
	for section, items in SECTIONS:
		for sheet, *_x in items:
			if sheet in results:
				c = ws.cell(row, 1, sheet)
				c.hyperlink = f"#'{sheet}'!A1"
				c.font = Font(name=F, size=10, color=TEAL, underline="single")
				ws.cell(row, 2, section).font = Font(name=F, size=10, color="6B6B6B")
				row += 1
	if errors:
		ws.cell(row + 1, 1, "Not included (see Error Log): " + ", ".join(errors)).font = Font(name=F, size=9, color="C0392B")
	buf = io.BytesIO()
	wb.save(buf)
	return buf.getvalue(), results


def add_chart(ws, chart, row, col):
	"""Write a report's chart data next to the table and draw a native Excel chart from it."""
	from openpyxl.chart import BarChart, LineChart, PieChart, Reference

	if not chart or not chart.get("data") or not chart["data"].get("labels"):
		return
	labels = chart["data"]["labels"]
	sets = chart["data"].get("datasets") or []
	if not sets:
		return
	from openpyxl.styles import Font
	ws.cell(row, col, "Chart data").font = Font(name="Arial", bold=True)
	for i, d in enumerate(sets):
		ws.cell(row, col + 1 + i, d.get("name") or "Value")
	for j, lab in enumerate(labels):
		ws.cell(row + 1 + j, col, str(lab))
		for i, d in enumerate(sets):
			vals = d.get("values") or []
			ws.cell(row + 1 + j, col + 1 + i, vals[j] if j < len(vals) else None)
	kind = chart.get("type")
	ch = LineChart() if kind == "line" else PieChart() if kind in ("donut", "pie") else BarChart()
	if isinstance(ch, BarChart) and len(labels) > 6:
		ch.type = "bar"
		ch.x_axis.scaling.orientation = "maxMin"
	ch.height, ch.width = 9, 18
	ch.add_data(Reference(ws, min_col=col + 1, max_col=col + len(sets), min_row=row, max_row=row + len(labels)),
		titles_from_data=True)
	ch.set_categories(Reference(ws, min_col=col, min_row=row + 1, max_row=row + len(labels)))
	palette = ["1F5F3F", "E9A23B", "2E8B6E", "E07A5F", "C0392B"]
	if not isinstance(ch, PieChart):
		for i, sr in enumerate(ch.series):
			if kind == "line":
				sr.graphicalProperties.line.solidFill = palette[i % 5]
			else:
				sr.graphicalProperties.solidFill = palette[i % 5]
		if len(sets) == 1:
			ch.legend = None
	ws.add_chart(ch, ws.cell(3, col).coordinate)


# ---------------------------------------------------------------- records, email, schedule

def attach_pack(doc):
	content, _r = build_workbook(doc.company, getdate(doc.period_start), getdate(doc.period_end), doc.commentary)
	fname = f"Management Pack {doc.company} {getdate(doc.period_start).strftime('%Y-%m')}.xlsx"
	for old in frappe.get_all("File", filters={"attached_to_doctype": doc.doctype, "attached_to_name": doc.name}, pluck="name"):
		frappe.delete_doc("File", old, ignore_permissions=True, force=True)
	f = frappe.get_doc({"doctype": "File", "file_name": fname, "attached_to_doctype": doc.doctype,
		"attached_to_name": doc.name, "attached_to_field": "pack_file", "is_private": 1, "content": content})
	f.insert(ignore_permissions=True)
	doc.db_set({"pack_file": f.file_url, "generated_on": now_datetime()})
	return content, fname


def recipients(s):
	listed = [x.strip() for x in (s.get("pack_recipients") or "").replace(",", "\n").splitlines() if x.strip()]
	if listed:
		return listed
	return frappe.db.sql_list("""select distinct u.name from `tabUser` u join `tabHas Role` r on r.parent = u.name
		where r.role = 'Management' and u.enabled = 1 and u.name not in ('Administrator', 'Guest')""")


def send_pack(doc):
	frappe.only_for(["Management", "System Manager"])
	from sales_board.utils.data import settings
	s = settings()
	to = recipients(s)
	if not to:
		frappe.throw("No recipients: add addresses under Monthly Management Pack in Sales Board Settings, or give users the Management role.")
	content, fname = attach_pack(doc)
	month = getdate(doc.period_start).strftime("%B %Y")
	body = f"<p>Please find attached the {month} management pack for {doc.company}.</p>"
	if doc.commentary:
		body += f"<h4>Commentary</h4>{doc.commentary}"
	frappe.sendmail(recipients=to, subject=f"{doc.company}: {month} management pack", message=body,
		attachments=[{"fname": fname, "fcontent": content}], reference_doctype=doc.doctype, reference_name=doc.name,
		now=True)
	doc.db_set({"status": "Sent", "sent_on": now_datetime(), "sent_to": "\n".join(to)})
	return ", ".join(to)


def generate_month(period_end=None, send=None):
	"""Create (or rebuild) the pack for the month ending period_end for every company with CRM orders.

	bench --site <site> execute sales_board.pack.generate_month --kwargs "{'period_end': '2026-09-30'}"
	"""
	from sales_board.utils.data import settings
	s = settings()
	end = getdate(period_end) if period_end else add_days(get_first_day(getdate(today())), -1)
	start, end = get_first_day(end), get_last_day(end)
	auto = cint(s.get("pack_auto_send")) if send is None else cint(send)
	made = []
	for company in frappe.get_all("Company", pluck="name"):
		if not frappe.db.count("CRM Order", {"company": company}):
			continue
		name = frappe.db.get_value("Sales Board Pack", {"company": company, "period_start": start})
		doc = frappe.get_doc("Sales Board Pack", name) if name else frappe.get_doc({"doctype": "Sales Board Pack",
			"company": company, "period_start": start, "period_end": end}).insert(ignore_permissions=True)
		if doc.status == "Sent":
			continue
		if auto:
			send_pack(doc)
		else:
			attach_pack(doc)
			notify_reviewer(doc, s)
		made.append(doc.name)
	frappe.db.commit()
	return made


def notify_reviewer(doc, s):
	reviewer = s.get("pack_reviewer")
	if not reviewer:
		return
	link = frappe.utils.get_url_to_form(doc.doctype, doc.name)
	month = getdate(doc.period_start).strftime("%B %Y")
	frappe.sendmail(recipients=[reviewer], subject=f"{doc.company}: {month} pack ready for review",
		message=f"<p>The {month} management pack is ready. Add your commentary and send it from "
			f"<a href='{link}'>{doc.name}</a>.</p>", reference_doctype=doc.doctype, reference_name=doc.name)
